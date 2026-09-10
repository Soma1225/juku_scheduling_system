"""Read-only parsing of an ``info.xlsx`` instructor snapshot and DB import logic."""

from collections import Counter
from io import BytesIO
import re
import sqlite3
import unicodedata

import openpyxl


SHEET_NAME = "講師情報"

_HEADER_ALIASES = {
    "external_instructor_id": {"講師番号", "教員番号"},
    "full_name": {"講師名漢字", "講師氏名", "氏名漢字", "氏名"},
    "short_name": {"講師名略", "講師略称", "略称"},
    "kana": {"講師名かな", "講師名カナ", "氏名かな", "氏名カナ"},
    "source_status": {"状態", "ステータス"},
}


def _normalize_header(value) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"[\s()（）:：・_\-]+", "", text)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return unicodedata.normalize("NFKC", str(value)).strip()


def _find_header(ws) -> tuple[int, dict[str, int]]:
    aliases = {
        field: {_normalize_header(alias) for alias in labels}
        for field, labels in _HEADER_ALIASES.items()
    }
    for row_number in range(1, min(ws.max_row, 100) + 1):
        found: dict[str, int] = {}
        for cell in ws[row_number]:
            normalized = _normalize_header(cell.value)
            for field, labels in aliases.items():
                if normalized in labels and field not in found:
                    found[field] = cell.column
        if len(found) == len(_HEADER_ALIASES):
            return row_number, found
    required = "、".join(["講師番号", "講師名（漢字）", "講師名（略）", "講師名（かな）", "状態"])
    raise ValueError(f"必要な見出しが見つかりません: {required}")


def parse_instructor_snapshot(source) -> list[dict]:
    """講師情報シートから許可された5項目だけを読み取る。"""
    workbook_source = BytesIO(source) if isinstance(source, bytes) else source
    workbook = openpyxl.load_workbook(workbook_source, data_only=True, read_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError("「講師情報」というシートが見つかりません")
        ws = workbook[SHEET_NAME]
        header_row, columns = _find_header(ws)
        records = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            record = {
                field: _cell_text(ws.cell(row=row_number, column=column).value)
                for field, column in columns.items()
            }
            if not any(record.values()):
                continue
            record["source_row"] = row_number
            records.append(record)
        return records
    finally:
        workbook.close()


def _split_name(value: str) -> tuple[str, str]:
    parts = re.split(r"[\s\u3000]+", value.strip(), maxsplit=1)
    return (parts[0], parts[1]) if len(parts) == 2 else (value.strip(), "")


def _validate_active_records(records: list[dict]) -> list[dict]:
    active = [record for record in records if record["source_status"].strip() != "退職"]
    labels = {
        "external_instructor_id": "講師番号",
        "full_name": "講師名（漢字）",
        "short_name": "講師名（略）",
        "kana": "講師名（かな）",
    }
    errors = []
    for record in active:
        missing = [label for field, label in labels.items() if not record[field]]
        if missing:
            errors.append(f"{record['source_row']}行目: {', '.join(missing)}が空です")
    for field, label in (("external_instructor_id", "講師番号"), ("short_name", "講師名（略）")):
        duplicates = sorted(value for value, count in Counter(r[field] for r in active if r[field]).items() if count > 1)
        if duplicates:
            errors.append(f"{label}が重複しています: {', '.join(duplicates)}")
    if errors:
        raise ValueError(" / ".join(errors))
    return active


def import_instructor_snapshot(conn: sqlite3.Connection, source) -> dict:
    """在籍扱いの講師だけを講師番号でupsertし、件数を返す。"""
    records = parse_instructor_snapshot(source)
    if not records:
        raise ValueError("講師データが1件も見つかりません")
    active = _validate_active_records(records)
    retired_count = len(records) - len(active)

    existing_by_number = {
        row[0]: row[1]
        for row in conn.execute(
            "SELECT external_instructor_id, instructor_id FROM INSTRUCTORS "
            "WHERE external_instructor_id IS NOT NULL"
        ).fetchall()
    }
    short_name_owners = {
        row[0]: row[1]
        for row in conn.execute(
            "SELECT short_name, instructor_id FROM INSTRUCTORS WHERE short_name IS NOT NULL"
        ).fetchall()
    }
    for record in active:
        owner = short_name_owners.get(record["short_name"])
        expected = existing_by_number.get(record["external_instructor_id"])
        if owner is not None and owner != expected:
            raise ValueError(
                f"講師名（略）「{record['short_name']}」は別の講師番号で使用されています"
            )

    created = 0
    updated = 0
    conn.execute("SAVEPOINT instructor_snapshot_import")
    try:
        for record in active:
            last_name, first_name = _split_name(record["full_name"])
            last_name_kana, first_name_kana = _split_name(record["kana"])
            instructor_id = existing_by_number.get(record["external_instructor_id"])
            if instructor_id is None:
                conn.execute(
                    """INSERT INTO INSTRUCTORS (
                           last_name, first_name, last_name_kana, first_name_kana,
                           external_instructor_id, short_name, academic_year, status
                       ) VALUES (?, ?, ?, ?, ?, ?, NULL, '在籍')""",
                    (last_name, first_name, last_name_kana, first_name_kana,
                     record["external_instructor_id"], record["short_name"]),
                )
                created += 1
            else:
                conn.execute(
                    """UPDATE INSTRUCTORS
                       SET last_name=?, first_name=?, last_name_kana=?, first_name_kana=?,
                           short_name=?, status='在籍'
                       WHERE instructor_id=?""",
                    (last_name, first_name, last_name_kana, first_name_kana,
                     record["short_name"], instructor_id),
                )
                updated += 1
        conn.execute("RELEASE SAVEPOINT instructor_snapshot_import")
        conn.commit()
    except Exception:
        conn.execute("ROLLBACK TO SAVEPOINT instructor_snapshot_import")
        conn.execute("RELEASE SAVEPOINT instructor_snapshot_import")
        raise

    return {
        "source_count": len(records),
        "created_count": created,
        "updated_count": updated,
        "retired_skipped_count": retired_count,
    }
