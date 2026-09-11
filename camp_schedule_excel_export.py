"""講習会の確定済み時間割を現行の25人5コマ形式Excelへ出力する。"""

from __future__ import annotations

import datetime
import io
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from db import format_grade_label, get_current_academic_fiscal_year, get_grade_at_fiscal_year


TEMPLATE_PATH = Path(__file__).with_name("assets") / "printout_25nin5koma_template.xlsx"
DAYS_PER_SHEET = 7
MAX_SHEETS = 4
MAX_DAYS = DAYS_PER_SHEET * MAX_SHEETS
INSTRUCTORS_PER_PERIOD = 25
PERIOD_COUNT = 5
FIRST_DATA_ROW = 4
ROWS_PER_PERIOD = 25
DAY_FIRST_COLUMNS = (2, 9, 16, 23, 30, 37, 44)  # B, I, P, W, AD, AK, AR
WEEKDAYS_JP = ("月", "火", "水", "木", "金", "土", "日")


class CampScheduleExportError(ValueError):
    """Excel出力を安全に作成できない場合のエラー。"""


class CampScheduleCapacityError(CampScheduleExportError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


@dataclass(frozen=True)
class AssignmentRow:
    session_date: datetime.date
    period_number: int
    session_id: int
    instructor_id: int
    instructor_name: str
    student_id: int
    student_name: str
    enrollment_year: int
    base_grade: int
    subject_name: str


def _load_rows(conn, camp_id: int) -> tuple[str, list[AssignmentRow]]:
    camp = conn.execute("SELECT camp_name FROM CAMPS WHERE camp_id=?", (camp_id,)).fetchone()
    if not camp:
        raise CampScheduleExportError("対象の講習会が見つかりません")
    rows = conn.execute(
        """
        SELECT ts.session_date,ts.period_number,se.session_id,i.instructor_id,
               COALESCE(NULLIF(i.short_name,''),i.last_name),
               st.student_id,st.last_name||st.first_name,st.enrollment_year,st.base_grade,
               sub.subject_name
        FROM SESSIONS se
        JOIN TIME_SLOTS ts ON ts.slot_id=se.slot_id
        JOIN INSTRUCTORS i ON i.instructor_id=se.instructor_id
        JOIN ASSIGNMENTS a ON a.session_id=se.session_id
        JOIN STUDENTS st ON st.student_id=a.student_id
        JOIN SUBJECTS sub ON sub.subject_id=a.subject_id
        WHERE se.camp_id=?
        ORDER BY ts.session_date,ts.period_number,i.instructor_id,se.session_id,a.assignment_id
        """,
        (camp_id,),
    ).fetchall()
    return camp[0], [
        AssignmentRow(datetime.date.fromisoformat(row[0]), *row[1:]) for row in rows
    ]


def validate_camp_schedule_export(conn, camp_id: int) -> tuple[str, list[AssignmentRow], list[str]]:
    """出力データを読み、欠落出力につながる問題をすべて返す。"""
    camp_name, rows = _load_rows(conn, camp_id)
    if not rows:
        return camp_name, rows, ["確定済みの講習会割当がありません"]

    errors: list[str] = []
    dates = sorted({row.session_date for row in rows})
    if len(dates) > MAX_DAYS:
        errors.append(f"授業日が{len(dates)}日あります。出力上限は{MAX_DAYS}日です")

    sessions: dict[int, list[AssignmentRow]] = defaultdict(list)
    period_sessions: dict[tuple[datetime.date, int], set[int]] = defaultdict(set)
    for row in rows:
        if not 1 <= row.period_number <= PERIOD_COUNT:
            errors.append(
                f"{row.session_date.isoformat()}の時限{row.period_number}はテンプレートの1〜{PERIOD_COUNT}限に収まりません"
            )
        sessions[row.session_id].append(row)
        period_sessions[(row.session_date, row.period_number)].add(row.session_id)

    for session_id, assignments in sessions.items():
        if len(assignments) > 2:
            first = assignments[0]
            errors.append(
                f"{first.session_date.isoformat()} {first.period_number}限のセッションID={session_id}に"
                f"生徒が{len(assignments)}人います（上限2人）"
            )
    for (session_date, period), session_ids in sorted(period_sessions.items()):
        if len(session_ids) > INSTRUCTORS_PER_PERIOD:
            errors.append(
                f"{session_date.isoformat()} {period}限は講師{len(session_ids)}人で、"
                f"テンプレート上限{INSTRUCTORS_PER_PERIOD}人を超えています"
            )
    return camp_name, rows, list(dict.fromkeys(errors))


def _date_heading(value: datetime.date) -> str:
    return f"{value.year}年{value.month}月{value.day}日({WEEKDAYS_JP[value.weekday()]})"


def _fill_period_times(sheet, periods: dict[int, tuple[str, str]]) -> None:
    for period in range(1, PERIOD_COUNT + 1):
        start_time, end_time = periods.get(period, ("", ""))
        row = FIRST_DATA_ROW + (period - 1) * ROWS_PER_PERIOD
        start_parts = start_time.split(":", 1) if ":" in start_time else ("", "")
        end_parts = end_time.split(":", 1) if ":" in end_time else ("", "")
        sheet.cell(row=row + 5, column=1, value=start_parts[0])
        sheet.cell(row=row + 6, column=1, value="：")
        sheet.cell(row=row + 7, column=1, value=start_parts[1])
        sheet.cell(row=row + 9, column=1, value="～")
        sheet.cell(row=row + 13, column=1, value=end_parts[0])
        sheet.cell(row=row + 14, column=1, value="：")
        sheet.cell(row=row + 15, column=1, value=end_parts[1])


def build_camp_schedule_workbook(conn, camp_id: int):
    """テンプレートを書式ごと複製し、確定済み割当だけを埋めたWorkbookを返す。"""
    camp_name, rows, errors = validate_camp_schedule_export(conn, camp_id)
    if errors:
        raise CampScheduleCapacityError(errors)
    if not TEMPLATE_PATH.is_file():
        raise CampScheduleExportError(f"Excelテンプレートが見つかりません: {TEMPLATE_PATH}")

    dates = sorted({row.session_date for row in rows})
    sheet_count = math.ceil(len(dates) / DAYS_PER_SHEET)
    workbook = openpyxl.load_workbook(TEMPLATE_PATH)
    template = workbook[workbook.sheetnames[0]]
    sheets = [template]
    for _ in range(1, sheet_count):
        sheets.append(workbook.copy_worksheet(template))
    for index, sheet in enumerate(sheets, start=1):
        sheet.title = "25人5コマ" if index == 1 else f"25人5コマ_{index}"

    periods = {
        row[0]: (row[1], row[2])
        for row in conn.execute(
            "SELECT period_number,start_time,end_time FROM PERIODS WHERE period_number BETWEEN 1 AND 5"
        ).fetchall()
    }
    grouped: dict[tuple[datetime.date, int], dict[int, list[AssignmentRow]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for row in rows:
        grouped[(row.session_date, row.period_number)][row.session_id].append(row)

    for sheet in sheets:
        _fill_period_times(sheet, periods)
    for date_index, session_date in enumerate(dates):
        sheet = sheets[date_index // DAYS_PER_SHEET]
        first_col = DAY_FIRST_COLUMNS[date_index % DAYS_PER_SHEET]
        sheet.cell(row=1, column=first_col, value=_date_heading(session_date))
        for period in range(1, PERIOD_COUNT + 1):
            sessions = grouped.get((session_date, period), {})
            ordered = sorted(
                sessions.values(),
                key=lambda items: (items[0].instructor_name, items[0].instructor_id, items[0].session_id),
            )
            for row_offset, assignments in enumerate(ordered):
                output_row = FIRST_DATA_ROW + (period - 1) * ROWS_PER_PERIOD + row_offset
                first = assignments[0]
                sheet.cell(output_row, first_col, first.instructor_name)
                for student_offset, assignment in enumerate(assignments):
                    fiscal_year = get_current_academic_fiscal_year(assignment.session_date)
                    grade = get_grade_at_fiscal_year(
                        assignment.enrollment_year, assignment.base_grade, fiscal_year
                    )
                    student_col = first_col + 1 + student_offset * 3
                    sheet.cell(output_row, student_col, format_grade_label(grade))
                    sheet.cell(output_row, student_col + 1, assignment.student_name)
                    sheet.cell(output_row, student_col + 2, assignment.subject_name)
    return workbook, camp_name


def export_camp_schedule_xlsx(conn, camp_id: int) -> tuple[bytes, str]:
    workbook, camp_name = build_camp_schedule_workbook(conn, camp_id)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), f"{camp_name}_時間割.xlsx"
