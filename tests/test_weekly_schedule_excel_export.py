import io
import sqlite3
import unittest
from pathlib import Path

import openpyxl
from timetable_snapshot_import import parse_timetable_snapshot

from weekly_schedule_excel_export import (
    TEMPLATE_PATH,
    WeeklyScheduleCapacityError,
    build_classroom_weekly_workbook,
    export_instructor_weekly_xlsx,
    export_student_weekly_xlsx,
)
import app
import layout
from core_migrations import ensure_core_schema
from page_closure_dates import add_closure_date, delete_closure_date


ROOT = Path(__file__).resolve().parents[1]


class WeeklyScheduleExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.execute("PRAGMA foreign_keys=ON")
        self.conn.executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
        self.conn.executemany(
            "INSERT INTO PERIODS(period_number,start_time,end_time) VALUES(?,?,?)",
            [(1, "14:20", "15:40"), (2, "15:50", "17:10"), (3, "17:20", "18:40"),
             (4, "19:00", "20:20"), (5, "20:30", "21:50")],
        )
        self.conn.executemany(
            """INSERT INTO STUDENTS
               (student_id,last_name,first_name,last_name_kana,first_name_kana,
                enrollment_year,base_grade,enrollment_status,gender)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            [
                (1, "山田", "花子", "やまだ", "はなこ", 2026, 8, "在籍", "女"),
                (2, "鈴木", "太郎", "すずき", "たろう", 2026, 7, "在籍", "男"),
            ],
        )
        self.conn.execute(
            """INSERT INTO INSTRUCTORS
               (instructor_id,last_name,first_name,last_name_kana,first_name_kana,short_name,academic_year,status)
               VALUES(1,'田中','一郎','たなか','いちろう','田中','B2','在籍')"""
        )
        self.conn.executemany(
            """INSERT INTO SUBJECTS
               (subject_id,course_category,grade_band,subject_group,subject_name)
               VALUES(?,?,?,?,?)""",
            [
                (1, "個別指導", "中学生", "数学", "数学"),
                (2, "戦略指導", "中学生", "教科フォロー", "教科フォロー(文系)"),
            ],
        )
        self.conn.execute(
            """INSERT INTO REGULAR_COURSE_ENROLLMENTS
               (student_id,subject_id,instructor_id,day_of_week,period_number,effective_start_date,effective_end_date)
               VALUES(1,1,1,'月',1,'2026-04-01',NULL)"""
        )
        self.conn.executemany(
            """INSERT INTO FOLLOW_COURSE_ENROLLMENTS
               (student_id,subject_id,instructor_id,day_of_week,period_number,effective_start_date,effective_end_date)
               VALUES(?,?,?,?,?,'2026-04-01',NULL)""",
            [(1, 2, 1, "火", 2), (2, 2, 1, "月", 1)],
        )
        self.conn.commit()

    def tearDown(self):
        self.conn.close()

    def test_student_export_uses_vba_calendar_coordinates_and_preserves_template(self):
        self.conn.executemany(
            "INSERT INTO CLOSURE_DATES(closure_date,closure_name) VALUES(?,?)",
            [("2026-09-14", "臨時休校"), ("2026-09-20", "秋季休校")],
        )
        original = openpyxl.load_workbook(TEMPLATE_PATH)
        original_merges = {str(item) for item in original["5コマ"].merged_cells.ranges}
        body, filename = export_student_weekly_xlsx(
            self.conn, 1, "2026-09-14", "2026-10-05"
        )
        workbook = openpyxl.load_workbook(io.BytesIO(body))
        sheet = workbook["生徒週間時間割"]

        self.assertEqual(filename, "山田花子_20260914-20261005_授業時間割.xlsx")
        self.assertEqual(workbook.sheetnames, ["生徒週間時間割"])
        self.assertEqual(sheet["B5"].value, "中2")
        self.assertEqual(sheet["I5"].value, "山田花子")
        self.assertEqual(sheet["S5"].value, "さん")
        self.assertEqual((sheet["G10"].value, sheet["G11"].value, sheet["G12"].value), (9, 14, "月"))
        self.assertEqual((sheet["H18"].value, sheet["H19"].value, sheet["H20"].value), (10, 1, "木"))
        self.assertEqual(sheet["G13"].value, "数学")
        self.assertNotIn("G13:G17", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["M13"].value, "秋季休校")
        self.assertIn("M13:M17", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["M13"].alignment.textRotation, 255)
        self.assertEqual(sheet["L21"].value, "数学")
        self.assertEqual(sheet.print_area, "'生徒週間時間割'!$A$1:$V$37")
        self.assertTrue(original_merges.issubset({str(item) for item in sheet.merged_cells.ranges}))

    def test_instructor_export_groups_two_students_in_one_slot(self):
        body, filename = export_instructor_weekly_xlsx(self.conn, 1)
        workbook = openpyxl.load_workbook(io.BytesIO(body))
        sheet = workbook["講師週間時間割"]

        self.assertEqual(filename, "田中一郎_講師週間時間割.xlsx")
        self.assertEqual(sheet["Q1"].value, "田中一郎")
        self.assertEqual(sheet["G3"].value, "14:20～15:40")
        self.assertEqual(sheet["W3"].value, "20:30～21:50")
        self.assertEqual(sheet["B4"].value, "月")
        self.assertEqual((sheet["H4"].value, sheet["I4"].value, sheet["J4"].value),
                         ("中2", "山田花子", "数学"))
        self.assertEqual((sheet["H5"].value, sheet["I5"].value, sheet["J5"].value),
                         ("中1", "鈴木太郎", "教科フォロー(文系)"))

    def test_classroom_export_matches_importable_column_layout(self):
        workbook = build_classroom_weekly_workbook(self.conn)
        sheet = workbook["時間割一覧"]

        self.assertEqual(sheet["B4"].value, "月")
        self.assertEqual((sheet["D4"].value, sheet["E4"].value, sheet["F4"].value, sheet["G4"].value),
                         ("田中", "c2 山田花子", "数学", None))
        self.assertEqual((sheet["D5"].value, sheet["E5"].value, sheet["F5"].value, sheet["G5"].value),
                         ("田中", "c1 鈴木太郎", "教科フォロー(文系)", None))
        self.assertEqual(sheet["B19"].value, "火")
        self.assertEqual(sheet.print_area, "'時間割一覧'!$B$1:$W$93")
        output = io.BytesIO()
        workbook.save(output)
        parsed = parse_timetable_snapshot(output.getvalue())
        self.assertEqual(len(parsed), 3)
        self.assertEqual(parsed[0]["student_text"], "c2 山田花子")

    def test_classroom_export_rejects_more_than_fifteen_rows(self):
        for student_id in range(3, 17):
            self.conn.execute(
                """INSERT INTO STUDENTS
                   (student_id,last_name,first_name,last_name_kana,first_name_kana,
                    enrollment_year,base_grade,enrollment_status)
                   VALUES(?,?,?,?,?,?,?,'在籍')""",
                (student_id, f"姓{student_id}", "名", f"せい{student_id}", "めい", 2026, 7),
            )
            self.conn.execute(
                """INSERT INTO REGULAR_COURSE_ENROLLMENTS
                   (student_id,subject_id,instructor_id,day_of_week,period_number,effective_start_date)
                   VALUES(?,1,1,'月',1,'2026-04-01')""",
                (student_id,),
            )
        with self.assertRaises(WeeklyScheduleCapacityError) as caught:
            build_classroom_weekly_workbook(self.conn)
        self.assertIn("上限15件", str(caught.exception))

    def test_student_export_rejects_overlapping_regular_and_follow(self):
        self.conn.execute(
            """INSERT INTO FOLLOW_COURSE_ENROLLMENTS
               (student_id,subject_id,instructor_id,day_of_week,period_number,effective_start_date)
               VALUES(1,2,1,'月',1,'2026-04-01')"""
        )
        with self.assertRaises(WeeklyScheduleCapacityError):
            export_student_weekly_xlsx(self.conn, 1, "2026-09-14", "2026-09-20")

    def test_student_export_rejects_more_than_thirty_two_days(self):
        with self.assertRaisesRegex(ValueError, "32日間以内"):
            export_student_weekly_xlsx(self.conn, 1, "2026-09-01", "2026-10-03")

    def test_closure_dates_can_be_added_and_deleted_but_not_duplicated(self):
        add_closure_date(self.conn, "2025-01-01", "年始休校")
        self.assertEqual(
            self.conn.execute("SELECT closure_name FROM CLOSURE_DATES WHERE closure_date='2025-01-01'").fetchone()[0],
            "年始休校",
        )
        with self.assertRaisesRegex(ValueError, "既に休校日"):
            add_closure_date(self.conn, "2025-01-01", "重複")
        delete_closure_date(self.conn, "2025-01-01")
        self.assertIsNone(
            self.conn.execute("SELECT 1 FROM CLOSURE_DATES WHERE closure_date='2025-01-01'").fetchone()
        )

    def test_existing_database_receives_closure_dates_migration_once(self):
        conn = sqlite3.connect(":memory:")
        conn.executescript(
            """CREATE TABLE APP_SCHEMA_MIGRATIONS(
                   migration_id TEXT PRIMARY KEY, applied_at TEXT NOT NULL
               );
               INSERT INTO APP_SCHEMA_MIGRATIONS VALUES('003_instructor_snapshot_import','2026-01-01');"""
        )
        self.assertTrue(ensure_core_schema(conn))
        self.assertIsNotNone(
            conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='CLOSURE_DATES'").fetchone()
        )
        self.assertFalse(ensure_core_schema(conn))
        conn.close()

    def test_export_page_is_routed_and_visible_in_sidebar(self):
        self.assertIn("/weekly-schedule-export", app.ROUTES)
        self.assertIn("/closure-dates", app.ROUTES)
        basic_menu = dict(next(items for name, items in layout.MENU_GROUPS if name == "基本設定"))
        self.assertEqual(basic_menu["/weekly-schedule-export"], "通常授業 週間Excel出力")
        self.assertEqual(basic_menu["/closure-dates"], "休校日設定")


if __name__ == "__main__":
    unittest.main()
