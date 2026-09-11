import io
import sqlite3
import unittest

import openpyxl

from camp_schedule_excel_export import (
    CampScheduleCapacityError,
    TEMPLATE_PATH,
    export_camp_schedule_xlsx,
)


class CampScheduleExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.executescript(
            """
            CREATE TABLE CAMPS(camp_id INTEGER PRIMARY KEY,camp_name TEXT);
            CREATE TABLE PERIODS(period_number INTEGER PRIMARY KEY,start_time TEXT,end_time TEXT);
            CREATE TABLE TIME_SLOTS(slot_id INTEGER PRIMARY KEY,session_date TEXT,period_number INTEGER);
            CREATE TABLE INSTRUCTORS(
              instructor_id INTEGER PRIMARY KEY,last_name TEXT,first_name TEXT,short_name TEXT
            );
            CREATE TABLE STUDENTS(
              student_id INTEGER PRIMARY KEY,last_name TEXT,first_name TEXT,
              enrollment_year INTEGER,base_grade INTEGER
            );
            CREATE TABLE SUBJECTS(subject_id INTEGER PRIMARY KEY,subject_name TEXT);
            CREATE TABLE SESSIONS(
              session_id INTEGER PRIMARY KEY,camp_id INTEGER,slot_id INTEGER,instructor_id INTEGER
            );
            CREATE TABLE ASSIGNMENTS(
              assignment_id INTEGER PRIMARY KEY,session_id INTEGER,student_id INTEGER,subject_id INTEGER
            );
            INSERT INTO CAMPS VALUES(1,'夏期講習会2026');
            INSERT INTO SUBJECTS VALUES(1,'数学');
            INSERT INTO PERIODS VALUES
              (1,'14:20','15:40'),(2,'15:50','17:10'),(3,'17:20','18:40'),
              (4,'19:00','20:20'),(5,'20:30','21:50');
            """
        )

    def tearDown(self):
        self.conn.close()

    def _add_assignment(
        self, *, assignment_id: int, date: str, period: int = 1,
        instructor_id: int | None = None, student_id: int | None = None,
        short_name: str | None = None, base_grade: int = 8,
    ):
        instructor_id = instructor_id or assignment_id
        student_id = student_id or assignment_id
        existing_slot = self.conn.execute(
            "SELECT slot_id FROM TIME_SLOTS WHERE session_date=? AND period_number=?",
            (date, period),
        ).fetchone()
        slot_id = existing_slot[0] if existing_slot else (
            self.conn.execute("SELECT COALESCE(MAX(slot_id),0)+1 FROM TIME_SLOTS").fetchone()[0]
        )
        session_id = assignment_id
        self.conn.execute(
            "INSERT OR IGNORE INTO TIME_SLOTS VALUES(?,?,?)", (slot_id, date, period)
        )
        self.conn.execute(
            "INSERT OR IGNORE INTO INSTRUCTORS VALUES(?,?,?,?)",
            (instructor_id, f"講師{instructor_id}", "太郎", short_name),
        )
        self.conn.execute(
            "INSERT OR IGNORE INTO STUDENTS VALUES(?,?,?,?,?)",
            (student_id, f"生徒{student_id}", "花子", 2026, base_grade),
        )
        self.conn.execute(
            "INSERT OR IGNORE INTO SESSIONS VALUES(?,?,?,?)",
            (session_id, 1, slot_id, instructor_id),
        )
        self.conn.execute(
            "INSERT INTO ASSIGNMENTS VALUES(?,?,?,1)",
            (assignment_id, session_id, student_id),
        )

    def test_exports_two_students_and_eight_days_to_two_formatted_sheets(self):
        self._add_assignment(
            assignment_id=1, date="2026-07-21", short_name="田",
        )
        # 同じセッションの2人目。
        self.conn.execute("INSERT INTO STUDENTS VALUES(2,'大阪','次郎',2026,7)")
        self.conn.execute("INSERT INTO ASSIGNMENTS VALUES(2,1,2,1)")
        for index in range(1, 8):
            self._add_assignment(
                assignment_id=index + 2,
                date=f"2026-07-{21 + index:02d}",
                instructor_id=index + 2,
                student_id=index + 2,
            )
        self.conn.commit()

        body, filename = export_camp_schedule_xlsx(self.conn, 1)
        workbook = openpyxl.load_workbook(io.BytesIO(body))
        template = openpyxl.load_workbook(TEMPLATE_PATH)
        self.assertEqual(filename, "夏期講習会2026_時間割.xlsx")
        self.assertEqual(workbook.sheetnames, ["25人5コマ", "25人5コマ_2"])
        first, second = workbook.worksheets
        self.assertEqual(first["B1"].value, "2026年7月21日(火)")
        self.assertEqual(first["AR1"].value, "2026年7月27日(月)")
        self.assertEqual(second["B1"].value, "2026年7月28日(火)")
        self.assertEqual(
            [first.cell(4, column).value for column in range(2, 9)],
            ["田", "中2", "生徒1花子", "数学", "中1", "大阪次郎", "数学"],
        )
        self.assertEqual((first["A9"].value, first["A11"].value), ("14", "20"))
        self.assertEqual((first["A17"].value, first["A19"].value), ("15", "40"))

        template_sheet = template.worksheets[0]
        self.assertEqual(
            {str(cell_range) for cell_range in first.merged_cells.ranges},
            {str(cell_range) for cell_range in template_sheet.merged_cells.ranges},
        )
        self.assertEqual(first["B4"].fill.fill_type, template_sheet["B4"].fill.fill_type)
        self.assertEqual(first["B4"].fill.fgColor.rgb, template_sheet["B4"].fill.fgColor.rgb)
        self.assertEqual(first["B4"].border.left.style, template_sheet["B4"].border.left.style)
        self.assertEqual(first.column_dimensions["D"].width, template_sheet.column_dimensions["D"].width)
        self.assertEqual(second["B4"].border.left.style, template_sheet["B4"].border.left.style)

    def test_more_than_25_instructors_in_one_period_aborts_export(self):
        for index in range(1, 27):
            self._add_assignment(assignment_id=index, date="2026-07-21")
        with self.assertRaises(CampScheduleCapacityError) as caught:
            export_camp_schedule_xlsx(self.conn, 1)
        self.assertIn("講師26人", str(caught.exception))
        self.assertIn("上限25人", str(caught.exception))

    def test_more_than_28_dates_aborts_export(self):
        import datetime

        start = datetime.date(2026, 7, 1)
        for index in range(29):
            self._add_assignment(
                assignment_id=index + 1,
                date=(start + datetime.timedelta(days=index)).isoformat(),
            )
        with self.assertRaises(CampScheduleCapacityError) as caught:
            export_camp_schedule_xlsx(self.conn, 1)
        self.assertIn("29日", str(caught.exception))
        self.assertIn("上限は28日", str(caught.exception))

    def test_28_dates_fit_exactly_in_four_sheets(self):
        import datetime

        start = datetime.date(2026, 7, 1)
        for index in range(28):
            self._add_assignment(
                assignment_id=index + 1,
                date=(start + datetime.timedelta(days=index)).isoformat(),
            )
        body, _ = export_camp_schedule_xlsx(self.conn, 1)
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["25人5コマ", "25人5コマ_2", "25人5コマ_3", "25人5コマ_4"],
        )
        self.assertEqual(workbook["25人5コマ_4"]["AR1"].value, "2026年7月28日(火)")


if __name__ == "__main__":
    unittest.main()
