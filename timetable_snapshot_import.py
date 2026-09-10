"""現行システムの「時間割一覧」スナップショットを一括取り込みする。"""

from __future__ import annotations

from datetime import date
from io import BytesIO
import re
import sqlite3

import openpyxl

from db import get_current_academic_fiscal_year, get_grade_at_fiscal_year
from excel_import import resolve_subject


SHEET_NAME = "時間割一覧"
WEEKDAYS = ("月", "火", "水", "木", "金", "土")
# period: (instructor column, student column, subject column)
PERIOD_COLUMNS = {
    1: (4, 5, 6),
    2: (8, 9, 10),
    3: (12, 13, 14),
    4: (16, 17, 18),
    5: (20, 21, 22),
}


class TimetableImportValidationError(ValueError):
    """自動確定できないセルが含まれる場合の検証エラー。"""

    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_student_text(value) -> dict:
    """「学年コード 氏名」を厳密に分離する。"""
    raw = _text(value)
    match = re.fullmatch(r"(s[1-6]|c[1-3]|k[1-3]|高卒)\s+(.+)", raw)
    if not match:
        raise ValueError("『学年コード 氏名』の形式ではありません")
    code = match.group(1)
    name = match.group(2).strip()
    if not name:
        raise ValueError("氏名が空です")
    if code.startswith("s"):
        grade = int(code[1:])
    elif code.startswith("c"):
        grade = 6 + int(code[1:])
    elif code.startswith("k"):
        grade = 9 + int(code[1:])
    else:
        grade = 13
    return {"grade_code": code, "base_grade": grade, "name": name}


def parse_timetable_snapshot(source) -> list[dict]:
    """ワークブックを変更せず、曜日ブロック（各15行）を読み取る。"""
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise TimetableImportValidationError([f"シート『{SHEET_NAME}』が見つかりません"])
        sheet = workbook[SHEET_NAME]
        records: list[dict] = []
        marker_rows: list[tuple[int, str]] = []
        for row in range(1, sheet.max_row + 1):
            marker = _text(sheet.cell(row, 2).value)
            if marker in WEEKDAYS:
                marker_rows.append((row, marker))

        seen_days: set[str] = set()
        for start_row, weekday in marker_rows:
            if weekday in seen_days:
                continue
            seen_days.add(weekday)
            last_instructor = {period: "" for period in PERIOD_COLUMNS}
            for row in range(start_row, start_row + 15):
                for period, (instructor_col, student_col, subject_col) in PERIOD_COLUMNS.items():
                    instructor = _text(sheet.cell(row, instructor_col).value)
                    if instructor:
                        last_instructor[period] = instructor
                    student = _text(sheet.cell(row, student_col).value)
                    if not student or student == "休校日":
                        continue
                    records.append({
                        "weekday": weekday,
                        "period": period,
                        "instructor_short_name": last_instructor[period],
                        "student_text": student,
                        "subject_text": _text(sheet.cell(row, subject_col).value),
                        "cell": sheet.cell(row, student_col).coordinate,
                    })
        return records
    finally:
        workbook.close()


def _split_name(name: str) -> tuple[str, str]:
    """単一の氏名文字列を既存の姓・名カラムへ安全に格納する。"""
    parts = name.split()
    if len(parts) >= 2:
        return parts[0], " ".join(parts[1:])
    return name, ""


def _stored_student_name(last_name: str, first_name: str) -> str:
    return f"{last_name or ''} {first_name or ''}".strip()


def _grade_to_code(grade: int) -> str:
    if 1 <= grade <= 6:
        return f"s{grade}"
    if 7 <= grade <= 9:
        return f"c{grade - 6}"
    if 10 <= grade <= 12:
        return f"k{grade - 9}"
    if grade == 13:
        return "高卒"
    return ""


def _resolve_records(conn: sqlite3.Connection, parsed: list[dict]) -> list[dict]:
    errors: list[str] = []
    resolved: list[dict] = []
    active_instructors = {
        row[1]: row[0]
        for row in conn.execute(
            "SELECT instructor_id, short_name FROM INSTRUCTORS "
            "WHERE short_name IS NOT NULL AND status <> '辞職'"
        )
    }
    occupied: dict[tuple[str, str, str, int], tuple[int, int]] = {}

    for record in parsed:
        location = f"{record['weekday']}曜 {record['period']}限 {record['cell']}"
        try:
            student = parse_student_text(record["student_text"])
        except ValueError as exc:
            errors.append(f"{location}: 生徒『{record['student_text']}』を判定できません（{exc}）")
            continue

        short_name = record["instructor_short_name"]
        instructor_id = active_instructors.get(short_name)
        if not short_name:
            errors.append(f"{location}: 教員略称が空です")
        elif instructor_id is None:
            errors.append(f"{location}: 教員略称『{short_name}』は講師マスタと完全一致しません")

        subject_text = record["subject_text"]
        if not subject_text:
            errors.append(f"{location}: 科目が空です")
            subject_id = None
        else:
            subject_result = resolve_subject(conn, subject_text, student["base_grade"], None)
            if subject_result["status"] == "matched":
                subject_id = subject_result["candidates"][0]["id"]
            elif subject_result["status"] == "ambiguous":
                labels = "、".join(candidate["label"] for candidate in subject_result["candidates"])
                errors.append(f"{location}: 科目『{subject_text}』を一意に決められません（候補: {labels}）")
                subject_id = None
            else:
                errors.append(f"{location}: 科目『{subject_text}』は科目マスタと一致しません")
                subject_id = None

        if instructor_id is None or subject_id is None:
            continue
        key = (student["grade_code"], student["name"], record["weekday"], record["period"])
        assignment = (instructor_id, subject_id)
        if key in occupied and occupied[key] != assignment:
            errors.append(f"{location}: 同じ生徒・曜日・限に異なる授業が重複しています")
            continue
        occupied[key] = assignment
        resolved.append({**record, **student, "instructor_id": instructor_id, "subject_id": subject_id})

    if errors:
        raise TimetableImportValidationError(errors)
    return resolved


def import_timetable_snapshot(
    conn: sqlite3.Connection,
    source,
    *,
    effective_start_date: str | None = None,
) -> dict[str, int]:
    """全セルを検証後、生徒と通常授業を同一トランザクションで反映する。"""
    start_date = effective_start_date or date.today().isoformat()
    try:
        date.fromisoformat(start_date)
    except ValueError as exc:
        raise ValueError("適用開始日は YYYY-MM-DD 形式で指定してください") from exc

    parsed = parse_timetable_snapshot(source)
    if not parsed:
        raise TimetableImportValidationError(["取り込み対象の授業が1件も見つかりません"])
    resolved = _resolve_records(conn, parsed)
    current_fy = get_current_academic_fiscal_year(date.fromisoformat(start_date))
    existing_students: dict[tuple[str, str], int] = {}
    for row in conn.execute(
        "SELECT student_id,last_name,first_name,enrollment_year,base_grade FROM STUDENTS "
        "ORDER BY student_id"
    ):
        current_grade = get_grade_at_fiscal_year(row[3], row[4], current_fy)
        key = (_grade_to_code(current_grade), _stored_student_name(row[1], row[2]))
        existing_students.setdefault(key, row[0])

    conn.execute("SAVEPOINT timetable_snapshot_import")
    try:
        new_students = 0
        new_enrollments = 0
        skipped_enrollments = 0
        for record in resolved:
            student_key = (record["grade_code"], record["name"])
            student_id = existing_students.get(student_key)
            if student_id is None:
                last_name, first_name = _split_name(record["name"])
                cur = conn.execute(
                    """INSERT INTO STUDENTS
                       (last_name,first_name,last_name_kana,first_name_kana,
                        enrollment_year,base_grade,enrollment_status)
                       VALUES (?,?,?,?,?,?,?)""",
                    (last_name, first_name, "", "", current_fy, record["base_grade"], "在籍"),
                )
                student_id = cur.lastrowid
                existing_students[student_key] = student_id
                new_students += 1

            exact = conn.execute(
                """SELECT 1 FROM REGULAR_COURSE_ENROLLMENTS
                   WHERE student_id=? AND subject_id=? AND instructor_id=?
                     AND day_of_week=? AND period_number=?
                     AND effective_start_date=? AND effective_end_date IS NULL""",
                (student_id, record["subject_id"], record["instructor_id"],
                 record["weekday"], record["period"], start_date),
            ).fetchone()
            if exact:
                skipped_enrollments += 1
                continue
            conflict = conn.execute(
                """SELECT 1 FROM REGULAR_COURSE_ENROLLMENTS
                   WHERE student_id=? AND day_of_week=? AND period_number=?
                     AND effective_start_date <= ?
                     AND (effective_end_date IS NULL OR effective_end_date >= ?)""",
                (student_id, record["weekday"], record["period"], start_date, start_date),
            ).fetchone()
            if conflict:
                raise TimetableImportValidationError([
                    f"{record['weekday']}曜 {record['period']}限: {record['name']}には既存の有効な通常授業があります"
                ])
            conn.execute(
                """INSERT INTO REGULAR_COURSE_ENROLLMENTS
                   (student_id,subject_id,instructor_id,day_of_week,period_number,
                    effective_start_date,effective_end_date)
                   VALUES (?,?,?,?,?,?,NULL)""",
                (student_id, record["subject_id"], record["instructor_id"],
                 record["weekday"], record["period"], start_date),
            )
            new_enrollments += 1
        conn.execute("RELEASE SAVEPOINT timetable_snapshot_import")
        conn.commit()
        return {
            "new_students": new_students,
            "new_enrollments": new_enrollments,
            "skipped_enrollments": skipped_enrollments,
        }
    except Exception:
        conn.execute("ROLLBACK TO SAVEPOINT timetable_snapshot_import")
        conn.execute("RELEASE SAVEPOINT timetable_snapshot_import")
        raise
