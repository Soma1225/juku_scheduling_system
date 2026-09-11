"""通常授業・教科フォローの週間時間割をExcelへ出力する。"""

from __future__ import annotations

import io
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from db import format_grade_label, get_current_grade


TEMPLATE_PATH = Path(__file__).with_name("assets") / "printout_student_teacher_weekly_template.xlsx"
WEEKDAYS = ("月", "火", "水", "木", "金", "土")
PERIOD_COUNT = 5
CLASSROOM_ROWS_PER_PERIOD = 15


class WeeklyScheduleExportError(ValueError):
    """週間時間割を安全に出力できない場合のエラー。"""


class WeeklyScheduleCapacityError(WeeklyScheduleExportError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


@dataclass(frozen=True)
class WeeklyLesson:
    student_id: int
    student_name: str
    enrollment_year: int
    base_grade: int
    instructor_id: int
    instructor_name: str
    subject_name: str
    day_of_week: str
    period_number: int
    course_type: str

    @property
    def current_grade(self) -> int:
        return get_current_grade(self.enrollment_year, self.base_grade)


def _load_lessons(conn, *, student_id: int | None = None, instructor_id: int | None = None,
                  active_people_only: bool = False) -> list[WeeklyLesson]:
    conditions = ["e.effective_end_date IS NULL", "e.day_of_week IN ('月','火','水','木','金','土')"]
    params: list[int] = []
    if student_id is not None:
        conditions.append("e.student_id = ?")
        params.append(student_id)
    if instructor_id is not None:
        conditions.append("e.instructor_id = ?")
        params.append(instructor_id)
    if active_people_only:
        conditions.extend(("st.enrollment_status = '在籍'", "i.status = '在籍'"))
    where = " AND ".join(conditions)
    sql = f"""
        SELECT e.student_id,st.last_name||st.first_name,st.enrollment_year,st.base_grade,
               e.instructor_id,COALESCE(NULLIF(i.short_name,''),i.last_name),
               sub.subject_name,e.day_of_week,e.period_number,?
        FROM {{table}} e
        JOIN STUDENTS st ON st.student_id=e.student_id
        JOIN INSTRUCTORS i ON i.instructor_id=e.instructor_id
        JOIN SUBJECTS sub ON sub.subject_id=e.subject_id
        WHERE {where}
    """
    rows = []
    for table, course_type in (
        ("REGULAR_COURSE_ENROLLMENTS", "通常授業"),
        ("FOLLOW_COURSE_ENROLLMENTS", "教科フォロー"),
    ):
        rows.extend(conn.execute(sql.format(table=table), [course_type, *params]).fetchall())
    day_order = {day: index for index, day in enumerate(WEEKDAYS)}
    rows.sort(key=lambda row: (day_order[row[7]], row[8], row[5], row[1], row[9]))
    return [WeeklyLesson(*row) for row in rows]


def _load_template():
    if not TEMPLATE_PATH.is_file():
        raise WeeklyScheduleExportError(f"Excelテンプレートが見つかりません: {TEMPLATE_PATH}")
    return openpyxl.load_workbook(TEMPLATE_PATH)


def _period_labels(conn) -> dict[int, str]:
    return {
        row[0]: f"{row[1]}～{row[2]}"
        for row in conn.execute(
            "SELECT period_number,start_time,end_time FROM PERIODS WHERE period_number BETWEEN 1 AND 5"
        )
    }


def _grade_code(grade: int) -> str:
    if 1 <= grade <= 6:
        return f"s{grade}"
    if 7 <= grade <= 9:
        return f"c{grade - 6}"
    if 10 <= grade <= 12:
        return f"k{grade - 9}"
    if grade == 13:
        return "高卒"
    raise WeeklyScheduleExportError(f"学年コードへ変換できない学年です: {grade}")


def _validate_slot_capacity(lessons: list[WeeklyLesson], *, limit: int, label: str) -> None:
    grouped: dict[tuple[str, int], list[WeeklyLesson]] = defaultdict(list)
    for lesson in lessons:
        if not 1 <= lesson.period_number <= PERIOD_COUNT:
            grouped[(lesson.day_of_week, lesson.period_number)].append(lesson)
        else:
            grouped[(lesson.day_of_week, lesson.period_number)].append(lesson)
    errors = []
    for (day, period), items in grouped.items():
        if not 1 <= period <= PERIOD_COUNT:
            errors.append(f"{day}曜 {period}限はテンプレートの1～5限に収まりません")
        elif len(items) > limit:
            errors.append(f"{day}曜 {period}限は{len(items)}件あり、{label}の上限{limit}件を超えています")
    if errors:
        raise WeeklyScheduleCapacityError(errors)


def build_student_weekly_workbook(conn, student_id: int):
    student = conn.execute(
        """SELECT last_name,first_name,enrollment_year,base_grade,gender
           FROM STUDENTS WHERE student_id=?""",
        (student_id,),
    ).fetchone()
    if student is None:
        raise WeeklyScheduleExportError("対象の生徒が見つかりません")
    lessons = _load_lessons(conn, student_id=student_id)
    _validate_slot_capacity(lessons, limit=1, label="生徒用時間割")

    workbook = _load_template()
    sheet = workbook["5コマ"]
    workbook.remove(workbook["講師6コマ"])
    name = f"{student[0]}{student[1]}"
    sheet["B5"] = format_grade_label(get_current_grade(student[2], student[3]))
    sheet["I5"] = name
    sheet["S5"] = "くん" if student[4] == "男" else "さん"

    period_labels = _period_labels(conn)
    sheet["B10"] = "科目"
    sheet["B18"] = "講師"
    sheet["B11"] = "曜日・時間帯"
    sheet["B19"] = "曜日・時間帯"
    for period in range(1, PERIOD_COUNT + 1):
        sheet.cell(12 + period, 2, period_labels.get(period, ""))
        sheet.cell(20 + period, 2, period_labels.get(period, ""))
    for day_index, day in enumerate(WEEKDAYS):
        column = 7 + day_index  # VBAの先頭日付列 G から月～土を配置
        sheet.cell(12, column, day)
        sheet.cell(20, column, day)
    for lesson in lessons:
        column = 7 + WEEKDAYS.index(lesson.day_of_week)
        for row, value in (
            (12 + lesson.period_number, lesson.subject_name),
            (20 + lesson.period_number, lesson.instructor_name),
        ):
            cell = sheet.cell(row, column)
            cell.value = value
            alignment = copy(cell.alignment)
            alignment.shrink_to_fit = True
            alignment.wrap_text = False
            alignment.horizontal = "center"
            alignment.vertical = "center"
            cell.alignment = alignment
    sheet.title = "生徒週間時間割"
    return workbook, name


def build_instructor_weekly_workbook(conn, instructor_id: int):
    instructor = conn.execute(
        "SELECT last_name,first_name FROM INSTRUCTORS WHERE instructor_id=?",
        (instructor_id,),
    ).fetchone()
    if instructor is None:
        raise WeeklyScheduleExportError("対象の講師が見つかりません")
    lessons = _load_lessons(conn, instructor_id=instructor_id)
    _validate_slot_capacity(lessons, limit=2, label="講師用時間割（1:2）")

    workbook = _load_template()
    sheet = workbook["講師6コマ"]
    workbook.remove(workbook["5コマ"])
    name = f"{instructor[0]}{instructor[1]}"
    sheet["Q1"] = name
    period_labels = _period_labels(conn)
    for period in range(1, PERIOD_COUNT + 1):
        # VBAの場所計算どおり、5限運用時は6ブロック中の後ろ5ブロックを使う。
        first_col = 7 + (period - 1) * 4  # G, K, O, S, W
        sheet.cell(3, first_col, period_labels.get(period, ""))
    grouped: dict[tuple[str, int], list[WeeklyLesson]] = defaultdict(list)
    for lesson in lessons:
        grouped[(lesson.day_of_week, lesson.period_number)].append(lesson)
    for day_index, day in enumerate(WEEKDAYS):
        first_row = 4 + day_index * 2
        sheet.cell(first_row, 2, day)
        for period in range(1, PERIOD_COUNT + 1):
            first_col = 7 + (period - 1) * 4
            for student_offset, lesson in enumerate(grouped.get((day, period), [])):
                row = first_row + student_offset
                sheet.cell(row, first_col + 1, format_grade_label(lesson.current_grade))
                sheet.cell(row, first_col + 2, lesson.student_name)
                sheet.cell(row, first_col + 3, lesson.subject_name)
                for column in range(first_col + 1, first_col + 4):
                    cell = sheet.cell(row, column)
                    alignment = copy(cell.alignment)
                    alignment.shrink_to_fit = True
                    alignment.wrap_text = False
                    alignment.horizontal = "center"
                    alignment.vertical = "center"
                    cell.alignment = alignment
    sheet.title = "講師週間時間割"
    return workbook, name


def build_classroom_weekly_workbook(conn):
    lessons = _load_lessons(conn, active_people_only=True)
    _validate_slot_capacity(lessons, limit=CLASSROOM_ROWS_PER_PERIOD, label="教室全体用時間割")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "時間割一覧"
    sheet["B1"] = "通常授業 週間時間割一覧"
    sheet["B1"].font = Font(size=16, bold=True, color="1F4E5F")
    sheet.merge_cells("B1:W1")
    sheet["B1"].alignment = Alignment(horizontal="center")

    dark = PatternFill("solid", fgColor="1F4E5F")
    pale = PatternFill("solid", fgColor="EAF1F3")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C2C5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ("教員", "生徒名", "科目", "出欠")
    for period in range(1, PERIOD_COUNT + 1):
        first_col = 4 + (period - 1) * 4
        sheet.merge_cells(start_row=2, start_column=first_col, end_row=2, end_column=first_col + 3)
        head = sheet.cell(2, first_col, f"{period}限")
        head.fill = dark
        head.font = white_bold
        head.alignment = Alignment(horizontal="center")
        for offset, label in enumerate(headers):
            cell = sheet.cell(3, first_col + offset, label)
            cell.fill = pale
            cell.font = Font(bold=True, color="1F4E5F")
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

    grouped: dict[tuple[str, int], list[WeeklyLesson]] = defaultdict(list)
    for lesson in lessons:
        grouped[(lesson.day_of_week, lesson.period_number)].append(lesson)
    for day_index, day in enumerate(WEEKDAYS):
        first_row = 4 + day_index * CLASSROOM_ROWS_PER_PERIOD
        sheet.cell(first_row, 2, day)
        sheet.cell(first_row, 2).fill = dark
        sheet.cell(first_row, 2).font = white_bold
        sheet.cell(first_row, 2).alignment = Alignment(horizontal="center")
        for row in range(first_row, first_row + CLASSROOM_ROWS_PER_PERIOD):
            for column in range(4, 24):
                sheet.cell(row, column).border = border
                sheet.cell(row, column).alignment = Alignment(horizontal="center", vertical="center")
        for period in range(1, PERIOD_COUNT + 1):
            first_col = 4 + (period - 1) * 4
            for offset, lesson in enumerate(grouped.get((day, period), [])):
                row = first_row + offset
                sheet.cell(row, first_col, lesson.instructor_name)
                sheet.cell(row, first_col + 1, f"{_grade_code(lesson.current_grade)} {lesson.student_name}")
                sheet.cell(row, first_col + 2, lesson.subject_name)

    widths = {"A": 2.5, "B": 5, "C": 2.5}
    for period in range(PERIOD_COUNT):
        first_col = 4 + period * 4
        widths.update({
            openpyxl.utils.get_column_letter(first_col): 9,
            openpyxl.utils.get_column_letter(first_col + 1): 17,
            openpyxl.utils.get_column_letter(first_col + 2): 11,
            openpyxl.utils.get_column_letter(first_col + 3): 6,
        })
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "D4"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = "B1:W93"
    return workbook


def _save(workbook, filename: str) -> tuple[bytes, str]:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), filename


def export_student_weekly_xlsx(conn, student_id: int) -> tuple[bytes, str]:
    workbook, name = build_student_weekly_workbook(conn, student_id)
    return _save(workbook, f"{name}_週間時間割.xlsx")


def export_instructor_weekly_xlsx(conn, instructor_id: int) -> tuple[bytes, str]:
    workbook, name = build_instructor_weekly_workbook(conn, instructor_id)
    return _save(workbook, f"{name}_講師週間時間割.xlsx")


def export_classroom_weekly_xlsx(conn) -> tuple[bytes, str]:
    return _save(build_classroom_weekly_workbook(conn), "教室全体_週間時間割.xlsx")
